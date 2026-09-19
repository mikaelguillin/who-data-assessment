import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from pipeline.adapters.base import AdapterResult, CountryAdapter
from pipeline.models import CountryAccount, HarmonisedRecord, QualityIssue
from pipeline.quality import (
    DATE_MAX_YEAR,
    DATE_MIN_YEAR,
    is_untrusted_description,
    labels_consistent,
    normalize_text,
)


def parse_amount_xof(raw: object) -> tuple[float | None, str | None]:
    if raw is None or str(raw).strip() == "":
        return None, "missing_amount"
    if isinstance(raw, int | float):
        return float(raw), None

    text = str(raw).strip().strip('"')
    cleaned = re.sub(r"(?i)fcfa", "", text)
    cleaned = cleaned.replace("\xa0", " ").replace(" ", "")
    if cleaned.count(",") == 1 and cleaned.count(".") == 0:
        left, right = cleaned.split(",")
        if right.isdigit() and len(right) <= 2 and left.replace("-", "").isdigit():
            cleaned = f"{left}.{right}"
        else:
            cleaned = cleaned.replace(",", "")
    else:
        cleaned = cleaned.replace(",", "")
    try:
        return float(cleaned), None
    except ValueError:
        return None, "unparseable_amount"


def parse_date_dmy(raw: object) -> datetime | None:
    if raw is None or str(raw).strip() == "":
        return None
    if isinstance(raw, datetime):
        return raw
    text = str(raw).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


class CountryBAdapter(CountryAdapter):
    source_format = "xlsx"
    layout_id = "xlsx_b"
    default_currency = "XOF"
    default_language = "fr"

    def load(self, path: Path, country_code: str) -> AdapterResult:
        filename = path.name
        workbook = load_workbook(path, data_only=True, read_only=True)
        official = self._load_plan_comptable(workbook)
        records = self._load_depenses(workbook, official, country_code, filename)
        workbook.close()

        accounts = [
            CountryAccount(
                country_code=country_code,
                account_code=code,
                account_label=label,
                source="official",
            )
            for code, label in official.items()
        ]
        return AdapterResult(
            country_code=country_code,
            source_filename=filename,
            source_format=self.source_format,
            layout_id=self.layout_id,
            records=records,
            accounts=accounts,
            notes="Excel extract with Plan_comptable sheet; preamble and TOTAL footer skipped",
            primary_currency=self.default_currency,
            language=self.default_language,
        )

    def _load_plan_comptable(self, workbook) -> dict[str, str]:
        sheet = workbook["Plan_comptable"]
        official: dict[str, str] = {}
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if not row or row[0] is None:
                continue
            official[str(row[0]).strip()] = str(row[1]).strip() if row[1] is not None else ""
        return official

    def _load_depenses(
        self,
        workbook,
        official: dict[str, str],
        country_code: str,
        filename: str,
    ) -> list[HarmonisedRecord]:
        sheet = workbook["Depenses"]
        header_index = None
        headers: list[str] = []
        records: list[HarmonisedRecord] = []

        for row_number, raw in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = list(raw)
            if header_index is None:
                first = str(values[0]).strip() if values and values[0] is not None else ""
                if first == "id_transaction":
                    header_index = row_number
                    headers = [str(v).strip() if v is not None else f"col_{i}" for i, v in enumerate(values)]
                continue

            first = values[0]
            if first is None or str(first).strip() == "":
                continue
            if str(first).strip().upper() == "TOTAL":
                continue

            row = {
                headers[i]: values[i] if i < len(values) else None for i in range(len(headers))
            }
            source_id = str(row.get("id_transaction") or "").strip()
            description = row.get("libelle")
            description_text = None if description is None else str(description)
            account_code = str(row.get("code_budgetaire") or "").strip() or None
            amount_raw = row.get("montant_XOF")
            amount, amount_flag = parse_amount_xof(amount_raw)
            parsed_date = parse_date_dmy(row.get("date_ecriture"))
            official_label = official.get(account_code) if account_code else None

            flags: list[QualityIssue] = []
            if amount_flag:
                amount_detail = None if amount_raw is None else str(amount_raw)
                flags.append(QualityIssue(flag_code=amount_flag, detail=amount_detail))
            if amount is not None and amount < 0:
                flags.append(QualityIssue(flag_code="negative_amount", detail=str(amount)))
            if parsed_date is None:
                flags.append(
                    QualityIssue(flag_code="unparseable_date", detail=str(row.get("date_ecriture")))
                )
            elif not DATE_MIN_YEAR <= parsed_date.year <= DATE_MAX_YEAR:
                flags.append(
                    QualityIssue(flag_code="date_out_of_range", detail=parsed_date.date().isoformat())
                )
            if not (description_text or "").strip():
                flags.append(QualityIssue(flag_code="missing_description"))
            if is_untrusted_description(description_text):
                flags.append(
                    QualityIssue(
                        flag_code="description_untrusted",
                        detail="Prompt-injection markers in libelle; classify from CoA only",
                    )
                )
            if (
                official_label
                and description_text
                and not labels_consistent(description_text, official_label)
            ):
                flags.append(
                    QualityIssue(
                        flag_code="coa_label_mismatch",
                        detail=f"libelle does not match Plan_comptable '{official_label}'",
                    )
                )
            if account_code and account_code not in official:
                flags.append(QualityIssue(flag_code="unknown_account_code", detail=account_code))

            payload = {key: (value.isoformat() if isinstance(value, datetime) else value) for key, value in row.items()}
            records.append(
                HarmonisedRecord(
                    country_code=country_code,
                    source_transaction_id=source_id,
                    source_row_ref=f"{filename}:Depenses:row:{row_number}",
                    transaction_date=parsed_date.date() if parsed_date else None,
                    fiscal_year=str(parsed_date.year) if parsed_date else None,
                    ministry_code=None if row.get("ministere_code") is None else str(row.get("ministere_code")),
                    ministry_name=None if row.get("ministere_nom") is None else str(row.get("ministere_nom")),
                    account_code=account_code,
                    description_raw=description_text,
                    description_norm=normalize_text(description_text),
                    supplier=None if row.get("tiers") is None else str(row.get("tiers")),
                    amount_original=None if amount_raw is None else str(amount_raw),
                    amount_native=amount,
                    currency_original="XOF",
                    payment_method=None,
                    raw_payload=payload,
                    official_account_label=official_label,
                    flags=flags,
                )
            )
        return records
